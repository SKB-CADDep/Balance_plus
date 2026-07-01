"""
Юнит-тест: Проверка правильности формирования Excel-отчета (матрицы, листы, округление).
"""
from openpyxl import load_workbook
from app.services.excel_exporter import ExcelExporter


class MockMatrixResult:
    def __init__(self, meta, columns, rows, values, warnings=None):
        self.meta = meta
        self.columns = columns
        self.rows = rows
        self.values = values
        self.warnings = warnings or []


class MockCalculationOutput:
    def __init__(self, tables, ejector_results):
        self.tables = tables
        self.ejector_results = ejector_results


def test_excel_structure_and_rounding():
    # 1. Подготавливаем фейковые входные данные (строго по структуре)
    tables = [
        MockMatrixResult(
            meta={"coefficient_b": 0.8, "W_main": 8000.0},
            columns=[100.0, 120.0],  # Колонки (G)
            rows=[15.0, 20.0],       # Строки (t1)
            values=[[10.12345, 0.00001],  # 0.00001 - проверка правила: не должен округляться до 0.0
                    [20.99999, 30.5]     # 20.99999 -> 21.0
                    ]
        )
    ]
    ejector_results = [
        {"P_ejector_kPa": 12.34567, "P_ejector_atm": 0.1218}
    ]
    calc_out = MockCalculationOutput(tables, ejector_results)

    # 2. Выполняем экспорт
    stream = ExcelExporter.export_calculation(calc_out)

    # 3. Открываем сгенерированный Excel
    wb = load_workbook(stream)

    # ПРОВЕРКА 1: Только один лист
    assert len(wb.sheetnames) == 1
    ws = wb.active
    assert ws.title == "Результаты расчёта"

    # ПРОВЕРКА 2: Мета-информация выведена корректно
    assert ws.cell(row=3, column=1).value == "coefficient_b"
    assert ws.cell(row=3, column=2).value == 0.8

    # ПРОВЕРКА 3: Матричный заголовок (t1 \ G)
    assert ws.cell(row=5, column=1).value == "t1 \\ G"
    assert ws.cell(row=5, column=2).value == 100.0  # G1
    assert ws.cell(row=5, column=3).value == 120.0  # G2

    # ПРОВЕРКА 4: Матричные значения и бизнес-правило округления
    assert ws.cell(row=6, column=1).value == 15.0  # t1_1
    assert ws.cell(row=6, column=2).value == 10.1235  # Округлилось до 4 знаков
    # Правило сработало: не стало нулем!
    assert ws.cell(row=6, column=3).value == 0.00001

    # ПРОВЕРКА 5: Наличие таблицы отсосов
    ejector_found = False
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == "Таблица по отсосам":
            ejector_found = True
            break
    assert ejector_found
