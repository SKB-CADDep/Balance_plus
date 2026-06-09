"""
Сервис генерации Excel-отчета по результатам расчёта конденсатора.

Принимает строго типизированный объект Pydantic (CalculationOutput),
декомпозирует процесс отрисовки на логические блоки и формирует
читаемый документ с автоматическим форматированием.
"""

import io
from typing import Any
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Font, PatternFill

# Импортируем нашу строгую Pydantic-схему (путь может немного отличаться в вашем проекте)
from app.schemas.base import CalculationOutput, MatrixResult, EjectorResult


class ExcelExporter:
    # =================================================================
    # КОНСТАНТЫ СТИЛЕЙ (Создаются один раз для экономии памяти)
    # =================================================================
    HEADER_FONT = Font(bold=True)
    META_FONT = Font(italic=True, color="404040")
    WARNING_FONT = Font(color="FF0000", italic=True)
    TABLE_HEADER_FILL = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")

    @staticmethod
    def _format_value(value: Any) -> Any:
        """
        Применяет бизнес-правила округления:
        - Округление до 4-х знаков.
        - Если округление даёт ровно 0.0, но исходное значение не ноль —
          возвращается полное значение (защита инженерных малых величин).
        """
        if isinstance(value, float):
            rounded = round(value, 4)
            if rounded == 0.0 and value != 0.0:
                return value
            return rounded
        return value

    @classmethod
    def export_calculation(cls, calc_output: CalculationOutput) -> io.BytesIO:
        """
        Главный метод-оркестратор генерации Excel-файла.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Результаты расчёта"

        current_row = 1

        # 1. Отрисовка матриц (таблиц режимов)
        if calc_output.tables:
            current_row = cls._write_matrices(ws, calc_output.tables, current_row)

        # 2. Отрисовка результатов по эжекторам (если есть)
        if calc_output.ejector_results:
            current_row = cls._write_ejectors(ws, calc_output.ejector_results, current_row)

        # 3. Финальное форматирование (ширина колонок)
        cls._adjust_column_widths(ws)

        # Сохранение в поток памяти
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        
        return stream

    @classmethod
    def _write_matrices(cls, ws: Worksheet, tables: list[MatrixResult], start_row: int) -> int:
        """Блок отрисовки двумерных матриц расчетов (участков)."""
        current_row = start_row

        for idx, table in enumerate(tables, 1):
            # Заголовок таблицы
            ws.cell(row=current_row, column=1, value=f"Таблица {idx} (Участок)").font = cls.HEADER_FONT
            current_row += 1

            # Метаданные (например, W и b)
            if table.meta:
                ws.cell(row=current_row, column=1, value="Режим:").font = cls.META_FONT
                current_row += 1
                for m_key, m_val in table.meta.items():
                    ws.cell(row=current_row, column=1, value=str(m_key)).font = cls.META_FONT
                    ws.cell(row=current_row, column=2, value=cls._format_value(m_val)).font = cls.META_FONT
                    current_row += 1

            # Заголовок осей: t1 (строки) \ G (столбцы)
            ws.cell(row=current_row, column=1, value="t1 \\ G").font = cls.HEADER_FONT
            ws.cell(row=current_row, column=1).fill = cls.TABLE_HEADER_FILL
            
            for c_idx, col_val in enumerate(table.columns, 2):
                cell = ws.cell(row=current_row, column=c_idx, value=cls._format_value(col_val))
                cell.font = cls.HEADER_FONT
                cell.fill = cls.TABLE_HEADER_FILL
            current_row += 1

            # Заполнение двумерного массива
            for r_idx, row_val in enumerate(table.rows):
                cell = ws.cell(row=current_row, column=1, value=cls._format_value(row_val))
                cell.font = cls.HEADER_FONT
                cell.fill = cls.TABLE_HEADER_FILL
                
                if r_idx < len(table.values):
                    for c_idx, val in enumerate(table.values[r_idx], 2):
                        ws.cell(row=current_row, column=c_idx, value=cls._format_value(val))
                current_row += 1
            
            # Отрисовка предупреждений
            if table.warnings:
                ws.cell(row=current_row, column=1, value="Предупреждения:").font = cls.WARNING_FONT
                current_row += 1
                for warning in table.warnings:
                    ws.cell(row=current_row, column=1, value=str(warning)).font = cls.WARNING_FONT
                    current_row += 1

            current_row += 2  # Отступ перед следующей таблицей

        return current_row

    @classmethod
    def _write_ejectors(cls, ws: Worksheet, ejectors: list[EjectorResult], start_row: int) -> int:
        """Блок отрисовки результатов работы воздухоудаляющих устройств."""
        current_row = start_row

        ws.cell(row=current_row, column=1, value="Таблица по отсосам").font = cls.HEADER_FONT
        current_row += 1

        # Берем ключи из первой Pydantic-модели для создания заголовков
        first_dict = ejectors[0].model_dump()
        headers = list(first_dict.keys())

        ws.cell(row=current_row, column=1, value="Отсос №").font = cls.HEADER_FONT
        ws.cell(row=current_row, column=1).fill = cls.TABLE_HEADER_FILL
        
        for c_idx, header in enumerate(headers, 2):
            cell = ws.cell(row=current_row, column=c_idx, value=str(header))
            cell.font = cls.HEADER_FONT
            cell.fill = cls.TABLE_HEADER_FILL
        current_row += 1

        # Заполняем данные
        for e_idx, ejector in enumerate(ejectors, 1):
            e_dict = ejector.model_dump()
            ws.cell(row=current_row, column=1, value=e_idx)
            for c_idx, header in enumerate(headers, 2):
                ws.cell(row=current_row, column=c_idx, value=cls._format_value(e_dict.get(header)))
            current_row += 1

        return current_row + 2

    @staticmethod
    def _adjust_column_widths(ws: Worksheet) -> None:
        """Автоматический подбор ширины колонок по содержимому."""
        for col in ws.columns:
            max_length = 0
            column_letter = col[0].column_letter
            for cell in col:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            # Ограничиваем ширину от 12 до 40 символов
            ws.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 40)